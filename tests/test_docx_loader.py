
import pytest
from docx import Document

from app.ingestion.docx_loader import load_docx

#Test that the DOCX loader extracts paragraphs correctly.
def test_load_docx_extracts_paragraphs(tmp_path):
    file_path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph("RAG retrieves relevant information.")
    document.save(file_path)

    result = load_docx(file_path)

    assert "RAG retrieves relevant information." in result

#Test that the DOCX loader extracts table text correctly.
def test_load_docx_extracts_table_text(tmp_path):
    file_path = tmp_path / "table.docx"

    document = Document()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Purpose"
    document.save(file_path)

    result = load_docx(file_path)

    assert "Component" in result
    assert "Purpose" in result

#Test that the DOCX loader raises a ValueError for unsupported file extensions.
def test_load_docx_rejects_unsupported_extension(tmp_path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text("Not a Word document", encoding="utf-8")

    with pytest.raises(ValueError):
        load_docx(file_path)

#Test that the DOCX loader raises a FileNotFoundError for missing files.
def test_load_docx_raises_for_missing_file(tmp_path):
    file_path = tmp_path / "missing.docx"

    with pytest.raises(FileNotFoundError):
        load_docx(file_path)
